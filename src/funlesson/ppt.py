"""大纲扩写为幻灯片文案（claude 单任务调用），再用 python-pptx 渲染成 pptx 文件。"""

import json
from dataclasses import asdict

from . import agent
from .models import Outline

_PROMPT_TEMPLATE = """你是一名教学助理，请把下面的课程大纲改写成适合做 PPT 的幻灯片文案。

要求：
- 每个一级大纲对应 1~2 页幻灯片，标题精炼，要点用简短的短语（不超过 20 字），每页 3~6 条要点
- 开头额外生成一页封面（title 用整体标题，bullets 为空数组）
- 只输出 JSON，不要输出任何解释性文字，格式：{{"slides": [{{"title": "...", "bullets": ["...", "..."]}}]}}

课程大纲（JSON）：
{outline_json}
"""


def build_ppt(outline: Outline, out_path: str) -> str:
    slides_data = _build_slides_content(outline)
    _render_pptx(slides_data, out_path)
    return out_path


def _build_slides_content(outline: Outline) -> list[dict]:
    prompt = _PROMPT_TEMPLATE.format(
        outline_json=json.dumps(asdict(outline), ensure_ascii=False)
    )
    data = agent.run_json_task(prompt, timeout=600)
    return data.get("slides", [])


def _render_pptx(slides_data: list[dict], out_path: str) -> None:
    from pptx import Presentation

    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]

    for i, slide_data in enumerate(slides_data):
        layout = title_layout if i == 0 else bullet_layout
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_data.get("title", "")

        bullets = slide_data.get("bullets") or []
        if bullets and len(slide.placeholders) > 1:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0]
            for b in bullets[1:]:
                p = body.add_paragraph()
                p.text = b

    prs.save(out_path)
